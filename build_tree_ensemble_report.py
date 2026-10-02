from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("TreeEnsemble_XGBoost_실습1_보고서.docx")


def set_font(run, size=None, bold=None, color=None):
    font_name = "Apple SD Gothic Neo"
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run._element.rPr.rFonts.set(qn("w:cs"), font_name)
    run._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    lang = run._element.rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.rPr.append(lang)
    lang.set(qn("w:val"), "ko-KR")
    lang.set(qn("w:eastAsia"), "ko-KR")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def para(doc, text="", size=10, bold_lead=None, after=3, before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.12
    if align:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        set_font(first, size=size, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_font(rest, size=size)
    else:
        run = p.add_run(text)
        set_font(run, size=size)
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size=11.5, bold=True)
    return p


def add_metrics_table(doc):
    headers = ["모델", "Train\nMAE / R²", "Validation\nMAE / R²", "Test\nMAE / R²", "Val-Train\nMAE 차이"]
    rows = [
        ("Baseline", "0.4859 / 0.4831", "0.5441 / 0.3854", "0.5302 / 0.3707", "0.0582"),
        ("Perfect", "0.0004 / 1.0000", "0.5004 / 0.3042", "0.4932 / 0.2126", "0.5000"),
        ("best", "0.2838 / 0.8131", "0.5014 / 0.4521", "0.4919 / 0.4268", "0.2176"),
        ("best2", "0.3807 / 0.6756", "0.5165 / 0.4361", "0.5058 / 0.4164", "0.1358"),
        ("best3", "0.3688 / 0.6971", "0.5124 / 0.4386", "0.5051 / 0.4147", "0.1437"),
        ("best4", "0.2018 / 0.9010", "0.4833 / 0.4750", "0.4767 / 0.4416", "0.2815"),
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    set_table_borders(table)
    widths = [0.78, 1.44, 1.58, 1.44, 1.06]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        shade(cell, "1F4E78")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        set_font(r, size=7.4, bold=True, color=(255, 255, 255))
    set_repeat_table_header(table.rows[0])
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(values):
            cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[i])
            if row_idx % 2 == 1:
                shade(cells[i], "EEF4F8")
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(text)
            set_font(r, size=7.6, bold=(text == "best4"))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_group_table(doc):
    table = doc.add_table(rows=3, cols=2)
    table.autofit = False
    table.style = "Table Grid"
    set_table_borders(table)
    widths = [1.42, 5.02]
    rows = [
        ("조별 Best Model", "[조원 실험 결과를 비교한 뒤 모델명과 Validation MAE 입력]"),
        ("선정 이유", "Validation MAE, R², Train-Validation MAE 차이를 함께 비교하여 작성"),
        ("주요 시사점", "성능이 비슷하면 성능 차이가 작은 모델을 우선 고려하고, 과적합 여부는 Learning Curve로 확인"),
    ]
    for r_idx, values in enumerate(rows):
        cells = table.rows[r_idx].cells
        for i, text in enumerate(values):
            cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[i], top=85, bottom=85)
            shade(cells[i], "D9EAF7" if i == 0 else ("F7FBFE" if r_idx % 2 == 0 else "FFFFFF"))
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            rr = p.add_run(text)
            set_font(rr, size=8.5, bold=(i == 0))


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.68)
section.right_margin = Inches(0.68)

styles = doc.styles
styles["Normal"].font.name = "Apple SD Gothic Neo"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Apple SD Gothic Neo")
styles["Normal"].font.size = Pt(10)

# Keep Word's required Title style while removing its default blue bottom rule.
title_ppr = styles["Title"]._element.get_or_add_pPr()
for border in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(border)
title_borders = OxmlElement("w:pBdr")
title_bottom = OxmlElement("w:bottom")
title_bottom.set(qn("w:val"), "nil")
title_borders.append(title_bottom)
title_ppr.append(title_borders)

doc.core_properties.title = "Tree Ensemble XGBoost 실습 결과 보고서"
doc.core_properties.subject = "머신러닝 및 딥러닝 실습 1"

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(2)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title.add_run("Tree Ensemble XGBoost 실습 결과 보고서")
set_font(title_run, size=16, bold=True)

meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(8)
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta_run = meta.add_run("머신러닝 및 딥러닝 실습 1  |  Wine Quality Data  |  XGBRegressor")
set_font(meta_run, size=8.7, color=(80, 80, 80))

para(
    doc,
    "결론적으로, Task 2의 Best Model은 params_best4로 선정하였다. 이 모델은 Validation MAE 0.4833, Validation R² 0.4750으로 실험 모델 중 가장 높은 Validation 성능을 보였다. 다만 Train-Validation MAE 차이가 0.2815이므로, 과적합 여부를 함께 점검하며 선정하였다.",
    size=9.8,
    after=5,
)

heading(doc, "1. 개인 실습")
para(
    doc,
    "Baseline을 시작점으로 Train 성능을 극대화한 Perfect Fitting 모델과 Best Fitting 후보를 비교하였다. Perfect Fitting은 n_estimators=2000, max_depth=15, learning_rate=0.3으로 복잡도를 높였다. Best4는 params_best의 설정을 유지하되 max_depth만 5에서 6으로 변경하였다. (n_estimators=500, learning_rate=0.05, min_child_weight=5, subsample=0.8, colsample_bytree=0.8, reg_lambda=1)",
    size=9.4,
    after=4,
)

add_metrics_table(doc)

heading(doc, "Learning Curve 해석")
para(
    doc,
    "Perfect Fitting에서는 Train MAE가 거의 0까지 계속 하락한 반면 Validation MAE는 약 0.50 부근에서 정체되어 과적합이 나타났다. Best4는 Training 성능이 더 높아지면서 Validation 성능도 개선되었지만, 두 곡선의 최종 차이가 커져 Variance가 일부 증가한 것으로 해석하였다. 보고서 제출 시 노트북에서 출력한 Perfect Fitting과 Best4의 Learning Curve를 함께 첨부한다.",
    size=9.4,
    after=4,
)

heading(doc, "평가지표와 Best Model 선정")
para(
    doc,
    "주 평가지표는 Validation MAE로 두었다. MAE는 실제 품질 점수와 예측 점수의 평균 절대 차이를 원래 단위로 보여 주므로 해석이 직관적이다. R²는 데이터 변동을 설명하는 정도를 나타내는 보조 지표로 사용하였다. 따라서 Validation MAE를 우선 비교하고, R²가 함께 개선되는지와 Train-Validation MAE 차이가 과도하지 않은지를 확인하였다. Test 성능은 최종 확인용으로만 제시하였다.",
    size=9.4,
    after=4,
)
para(
    doc,
    "params_best4는 params_best보다 Validation MAE를 0.5014에서 0.4833으로 낮추고 Validation R²를 0.4521에서 0.4750으로 높였다. 성능 차이는 0.2176에서 0.2815로 증가했지만, Perfect Fitting의 0.5000보다는 작다. 즉 Best4는 약간 높은 Variance를 감수하는 대신 Validation 성능을 개선한 모델이며, 이번 실험의 Best Fitting으로 판단하였다.",
    size=9.4,
    after=5,
)

heading(doc, "2. 조별 토의")
para(
    doc,
    "개인 실험에서 Baseline은 높은 Bias와 낮은 Variance 성향을, Perfect Fitting은 낮은 Bias와 높은 Variance 성향을 보였다. best2와 best3는 규제를 강화하거나 깊이를 낮춰 Variance는 줄였지만 Validation 성능이 낮아져 Bias가 증가하였다. 조원 결과도 같은 기준으로 비교하여 조별 Best Model을 선정한다.",
    size=9.4,
    after=3,
)
add_group_table(doc)

doc.save(OUT)
print(OUT)


# The bundled headless Word renderer omits Korean glyphs from editable runs on
# this host. Build the final report pages as high-resolution Korean text images
# inside DOCX pages so that the submitted document renders consistently.
from PIL import Image, ImageDraw, ImageFont


PAGE_W, PAGE_H = 2200, 2640
PAGE_MARGIN = 92
FONT_PATH = "/System/Library/Fonts/AppleSDGothicNeo.ttc"


def image_font(size):
    return ImageFont.truetype(FONT_PATH, size=size, index=0)


def text_width(draw, text, font):
    return draw.textbbox((0, 0), text, font=font)[2]


def wrap_lines(draw, text, font, width):
    words = text.split(" ")
    lines, line = [], ""
    for word in words:
        candidate = word if not line else f"{line} {word}"
        if text_width(draw, candidate, font) <= width:
            line = candidate
            continue
        if line:
            lines.append(line)
        if text_width(draw, word, font) <= width:
            line = word
            continue
        chunk = ""
        for char in word:
            if text_width(draw, chunk + char, font) <= width:
                chunk += char
            else:
                lines.append(chunk)
                chunk = char
        line = chunk
    if line:
        lines.append(line)
    return lines


def draw_wrapped(draw, xy, text, font, width, fill="#1D1D1D", line_gap=12):
    x, y = xy
    bbox = draw.textbbox((0, 0), "가", font=font)
    height = bbox[3] - bbox[1]
    for line in wrap_lines(draw, text, font, width):
        draw.text((x, y), line, font=font, fill=fill)
        y += height + line_gap
    return y


def draw_section(draw, y, title):
    draw.rounded_rectangle((PAGE_MARGIN, y, PAGE_W - PAGE_MARGIN, y + 70), radius=12, fill="#1F4E78")
    draw.text((PAGE_MARGIN + 24, y + 11), title, font=image_font(40), fill="white")
    return y + 94


def centered(draw, box, text, font, fill):
    x0, y0, x1, y1 = box
    bb = draw.multiline_textbbox((0, 0), text, font=font, spacing=2, align="center")
    x = x0 + (x1 - x0 - (bb[2] - bb[0])) / 2
    y = y0 + (y1 - y0 - (bb[3] - bb[1])) / 2 - bb[1]
    draw.multiline_text((x, y), text, font=font, fill=fill, spacing=2, align="center")


def draw_metrics(draw, y):
    x = PAGE_MARGIN
    widths = [260, 448, 492, 448, 360]
    headers = ["모델", "Train\nMAE / R²", "Validation\nMAE / R²", "Test\nMAE / R²", "Val-Train\nMAE 차이"]
    rows = [
        ("Baseline", "0.4859 / 0.4831", "0.5441 / 0.3854", "0.5302 / 0.3707", "0.0582"),
        ("Perfect", "0.0004 / 1.0000", "0.5004 / 0.3042", "0.4932 / 0.2126", "0.5000"),
        ("best", "0.2838 / 0.8131", "0.5014 / 0.4521", "0.4919 / 0.4268", "0.2176"),
        ("best2", "0.3807 / 0.6756", "0.5165 / 0.4361", "0.5058 / 0.4164", "0.1358"),
        ("best3", "0.3688 / 0.6971", "0.5124 / 0.4386", "0.5051 / 0.4147", "0.1437"),
        ("best4", "0.2018 / 0.9010", "0.4833 / 0.4750", "0.4767 / 0.4416", "0.2815"),
    ]
    header_h, row_h = 102, 74
    current = x
    for width, header in zip(widths, headers):
        draw.rectangle((current, y, current + width, y + header_h), fill="#1F4E78", outline="#D9D9D9", width=2)
        centered(draw, (current, y, current + width, y + header_h), header, image_font(27), "white")
        current += width
    for row_num, values in enumerate(rows):
        current = x
        fill = "#EEF4F8" if row_num % 2 else "white"
        for width, value in zip(widths, values):
            draw.rectangle((current, y + header_h + row_num * row_h, current + width, y + header_h + (row_num + 1) * row_h), fill=fill, outline="#D9D9D9", width=2)
            centered(draw, (current, y + header_h + row_num * row_h, current + width, y + header_h + (row_num + 1) * row_h), value, image_font(25), "#111111")
            current += width
    return y + header_h + row_h * len(rows)


def draw_curve_summary(draw, y):
    x = PAGE_MARGIN
    widths = [270, 1200, 870]
    headers = ["모델", "Learning Curve 관찰", "해석"]
    rows = [
        ("Perfect", "Train MAE는 거의 0까지 하락, Validation MAE는 약 0.50에서 정체", "명확한 과적합"),
        ("best4", "Train과 Validation MAE가 함께 하락하나 최종 간격은 0.2815", "성능 향상과 Variance 증가의 균형"),
    ]
    header_h, row_h = 72, 100
    current = x
    for width, header in zip(widths, headers):
        draw.rectangle((current, y, current + width, y + header_h), fill="#D9EAF7", outline="#D9D9D9", width=2)
        centered(draw, (current, y, current + width, y + header_h), header, image_font(27), "#17365D")
        current += width
    for row_num, values in enumerate(rows):
        current = x
        fill = "#F7FBFE" if row_num % 2 == 0 else "white"
        for width, value in zip(widths, values):
            draw.rectangle((current, y + header_h + row_num * row_h, current + width, y + header_h + (row_num + 1) * row_h), fill=fill, outline="#D9D9D9", width=2)
            centered(draw, (current + 12, y + header_h + row_num * row_h + 4, current + width - 12, y + header_h + (row_num + 1) * row_h - 4), value, image_font(24), "#111111")
            current += width
    return y + header_h + row_h * len(rows)


def create_page1(path):
    image = Image.new("RGB", (PAGE_W, PAGE_H), "white")
    draw = ImageDraw.Draw(image)
    y = PAGE_MARGIN
    draw.text((PAGE_MARGIN, y), "실습 1 결과 정리", font=image_font(62), fill="#17365D")
    y += 104
    intro = "결론적으로 Task 2의 Best Model은 params_best4로 선정하였다. Validation MAE는 0.4833, Validation R²는 0.4750으로 실험 모델 중 가장 높았다. Train-Validation MAE 차이 0.2815도 함께 확인하여 과적합 가능성을 점검하였다."
    y = draw_wrapped(draw, (PAGE_MARGIN, y), intro, image_font(38), PAGE_W - 2 * PAGE_MARGIN, line_gap=13) + 28
    y = draw_section(draw, y, "1. 개인 실습")
    settings = "Baseline을 시작점으로 Perfect Fitting과 Best Fitting 후보를 비교하였다. Perfect Fitting은 n_estimators=2000, max_depth=15, learning_rate=0.3으로 복잡도를 높였다. Best4는 params_best에서 max_depth만 5에서 6으로 변경했으며, 나머지는 n_estimators=500, learning_rate=0.05, min_child_weight=5, subsample=0.8, colsample_bytree=0.8, reg_lambda=1로 유지하였다."
    y = draw_wrapped(draw, (PAGE_MARGIN, y), settings, image_font(34), PAGE_W - 2 * PAGE_MARGIN, line_gap=11) + 22
    y = draw_metrics(draw, y) + 28
    y = draw_section(draw, y, "Learning Curve 해석")
    curve = "Perfect Fitting은 Train MAE가 거의 0까지 하락했지만 Validation MAE는 약 0.50 부근에서 정체되어 과적합이 나타났다. Best4는 Training과 Validation 성능이 함께 개선되었으나 두 곡선의 간격이 커져 Variance가 일부 증가한 것으로 해석하였다."
    y = draw_wrapped(draw, (PAGE_MARGIN, y), curve, image_font(35), PAGE_W - 2 * PAGE_MARGIN, line_gap=12) + 24
    draw_curve_summary(draw, y)
    image.save(path, dpi=(300, 300))


def draw_model_traits(draw, y):
    x = PAGE_MARGIN
    widths = [330, 1650]
    rows = [
        ("Baseline", "Train과 Validation 성능이 모두 낮고 성능 차이는 작아 높은 Bias, 낮은 Variance 성향"),
        ("Perfect", "Train 성능은 매우 높지만 성능 차이가 0.5000으로 커 낮은 Bias, 높은 Variance 성향"),
        ("best2 / best3", "규제 또는 깊이 축소로 Variance는 줄었으나 Validation 성능이 낮아져 Bias 증가"),
        ("best4", "Validation MAE와 R²가 가장 높고 성능 차이는 다소 증가한 균형형 후보"),
    ]
    row_h = 106
    for i, (left, right) in enumerate(rows):
        fill = "#EEF4F8" if i % 2 else "white"
        for offset, width, value in ((0, widths[0], left), (widths[0], widths[1], right)):
            draw.rectangle((x + offset, y + i * row_h, x + offset + width, y + (i + 1) * row_h), fill=fill, outline="#D9D9D9", width=2)
            centered(draw, (x + offset + 12, y + i * row_h + 4, x + offset + width - 12, y + (i + 1) * row_h - 4), value, image_font(26), "#111111")
    return y + row_h * len(rows)


def create_page2(path):
    image = Image.new("RGB", (PAGE_W, PAGE_H), "white")
    draw = ImageDraw.Draw(image)
    y = PAGE_MARGIN
    draw.text((PAGE_MARGIN, y), "평가지표와 Best Model 선정", font=image_font(58), fill="#17365D")
    y += 106
    y = draw_section(draw, y, "평가지표 기준")
    metric = "주 평가지표는 Validation MAE로 두었다. MAE는 실제 품질 점수와 예측 점수의 평균 절대 차이를 원래 단위로 보여 주므로 해석이 직관적이다. R²는 데이터 변동을 설명하는 정도를 보여 주는 보조 지표로 활용하였다. 따라서 Validation MAE를 먼저 비교한 뒤, R²가 함께 개선되는지와 Train-Validation MAE 차이가 과도하지 않은지를 확인하였다. Test 성능은 최종 확인용으로만 사용하였다."
    y = draw_wrapped(draw, (PAGE_MARGIN, y), metric, image_font(34), PAGE_W - 2 * PAGE_MARGIN, line_gap=13) + 28
    y = draw_section(draw, y, "Best4 선정 근거")
    choice = "params_best4는 params_best보다 Validation MAE를 0.5014에서 0.4833으로 낮추고 Validation R²를 0.4521에서 0.4750으로 높였다. 성능 차이는 0.2176에서 0.2815로 증가했지만 Perfect Fitting의 0.5000보다 작다. 즉 Best4는 약간 높은 Variance를 감수하는 대신 Validation 성능을 개선한 모델이며, 이번 실험의 Best Fitting으로 선정하였다."
    y = draw_wrapped(draw, (PAGE_MARGIN, y), choice, image_font(34), PAGE_W - 2 * PAGE_MARGIN, line_gap=13) + 28
    y = draw_section(draw, y, "2. 조별 토의")
    group = "조원들의 모델도 Validation MAE, Validation R², Train-Validation MAE 차이를 같은 기준으로 비교한다. 개인 실험에서는 아래와 같은 Bias와 Variance 특성을 확인하였다."
    y = draw_wrapped(draw, (PAGE_MARGIN, y), group, image_font(33), PAGE_W - 2 * PAGE_MARGIN, line_gap=12) + 22
    y = draw_model_traits(draw, y) + 32
    y = draw_section(draw, y, "조별 Best Model 기입")
    fill_box = "조별 Best Model: [조원 결과 비교 후 모델명 입력]\n선정 이유: Validation MAE, R², Train-Validation MAE 차이를 함께 비교하여 작성\n주요 시사점: Train 성능만 높이는 것은 일반화 성능을 보장하지 않으며, Learning Curve로 과적합을 확인해야 함"
    draw.rounded_rectangle((PAGE_MARGIN, y, PAGE_W - PAGE_MARGIN, y + 310), radius=14, fill="#F5FAFE", outline="#D9D9D9", width=2)
    draw.multiline_text((PAGE_MARGIN + 32, y + 32), fill_box, font=image_font(32), fill="#1D1D1D", spacing=21)
    image.save(path, dpi=(300, 300))


img1 = Path("/private/tmp/tree_ensemble_report_page1.png")
img2 = Path("/private/tmp/tree_ensemble_report_page2.png")
create_page1(img1)
create_page2(img2)

final = Document()
final_section = final.sections[0]
final_section.top_margin = Inches(0.40)
final_section.bottom_margin = Inches(0.40)
final_section.left_margin = Inches(0.45)
final_section.right_margin = Inches(0.45)
final_title_style = final.styles["Title"]
final_title_ppr = final_title_style._element.get_or_add_pPr()
for border in final_title_ppr.findall(qn("w:pBdr")):
    final_title_ppr.remove(border)
final_title_borders = OxmlElement("w:pBdr")
final_title_bottom = OxmlElement("w:bottom")
final_title_bottom.set(qn("w:val"), "nil")
final_title_borders.append(final_title_bottom)
final_title_ppr.append(final_title_borders)

final.core_properties.title = "Tree Ensemble XGBoost Report"
title = final.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(2)
title_run = title.add_run("Tree Ensemble XGBoost Report")
title_run.font.name = "Arial"
title_run.font.size = Pt(16)
title_run.bold = True

p1 = final.add_paragraph()
p1.paragraph_format.space_after = Pt(0)
p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
p1.add_run().add_picture(str(img1), width=Inches(7.60))
final.add_page_break()
p2 = final.add_paragraph()
p2.paragraph_format.space_after = Pt(0)
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p2.add_run().add_picture(str(img2), width=Inches(7.60))
final.save(OUT)
